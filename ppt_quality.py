import re
from dataclasses import dataclass, asdict
from typing import Dict, List, Tuple

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_AUTO_SIZE

_CJK_RE = re.compile(r"[\u3000-\u9fff\uac00-\ud7af\uf900-\ufaff]")
_HAN_RE = re.compile(r"[\u4e00-\u9fff]")
_JP_KANA_RE = re.compile(r"[\u3041-\u3096\u30a1-\u30fa\u30fd-\u30ff]")
_LATIN_RE = re.compile(r"[A-Za-z]")

DOUBLE_PATTERNS = [
    (re.compile(r"AI\s*AI", re.IGNORECASE), "AI重复"),
    (re.compile(r"DX\s*DX", re.IGNORECASE), "DX重复"),
    (re.compile(r"APP\s*APP", re.IGNORECASE), "APP重复"),
    (re.compile(r"※\s*※"), "※重复"),
    (re.compile(r"(\d{4})\1"), "4位数字重复(年份)"),
    (re.compile(r"([①②③④⑤⑥⑦⑧⑨⑩❶❷❸❹❺❻❼❽❾❿])\1"), "序号圆圈重复"),
]

TABLE_BULLET_PREFIX = re.compile(r"^[\s]*[•·●○◆▶→\-—]\s")
UNCHANGED_RATIO_THRESHOLD = 0.10
SOURCE_RESIDUAL_RATIO_THRESHOLD = 0.08


@dataclass
class PPTQAReport:
    passed: bool
    translation_accuracy_ok: bool
    format_layout_ok: bool
    text_color_ok: bool
    title_single_line_ok: bool
    summary: str
    issues: List[str]
    stats: Dict[str, int]
    source_lang_detected: str = "Unknown"
    target_lang: str = ""

    def to_dict(self) -> Dict:
        return asdict(self)


def _iter_shapes_recursive(shapes):
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):
            for child in _iter_shapes_recursive(shape.shapes):
                yield child


def _iter_paragraph_texts(prs: Presentation) -> List[Tuple[str, str, int]]:
    rows = []
    for slide_idx, slide in enumerate(prs.slides, 1):
        for shape in _iter_shapes_recursive(slide.shapes):
            if getattr(shape, "has_text_frame", False):
                for para in shape.text_frame.paragraphs:
                    txt = (para.text or "").strip()
                    if txt:
                        rows.append(("para", txt, slide_idx))
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            txt = (para.text or "").strip()
                            if txt:
                                rows.append(("cell", txt, slide_idx))
    return rows


def _detect_source_language(rows: List[Tuple[str, str, int]]) -> str:
    han_hits = 0
    kana_hits = 0
    latin_hits = 0

    for _kind, text, _slide_idx in rows:
        if _HAN_RE.search(text):
            han_hits += 1
        if _JP_KANA_RE.search(text):
            kana_hits += 1
        if _LATIN_RE.search(text):
            latin_hits += 1

    if han_hits > 0:
        return "Chinese"
    if kana_hits > 0:
        return "Japanese"
    if latin_hits > 0:
        return "English"
    return "Unknown"


def _color_signature(run) -> str:
    color = run.font.color
    if color is None:
        return "NONE"

    rgb = None
    try:
        if color.rgb is not None:
            rgb = str(color.rgb)
    except Exception:
        rgb = None

    theme = None
    try:
        if color.theme_color is not None:
            theme = str(color.theme_color)
    except Exception:
        theme = None

    brightness = None
    try:
        if color.brightness is not None:
            brightness = str(color.brightness)
    except Exception:
        brightness = None

    return f"rgb={rgb}|theme={theme}|brightness={brightness}"


def _collect_first_run_colors(prs: Presentation) -> List[str]:
    colors = []
    for slide in prs.slides:
        for shape in _iter_shapes_recursive(slide.shapes):
            if getattr(shape, "has_text_frame", False):
                for para in shape.text_frame.paragraphs:
                    if para.runs:
                        colors.append(_color_signature(para.runs[0]))
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            if para.runs:
                                colors.append(_color_signature(para.runs[0]))
    return colors


def enforce_title_single_line(output_path: str) -> Dict[str, int]:
    prs = Presentation(output_path)
    touched_titles = 0
    replaced_line_breaks = 0

    for slide in prs.slides:
        for shape in _iter_shapes_recursive(slide.shapes):
            if not getattr(shape, "has_text_frame", False):
                continue
            if not getattr(shape, "is_placeholder", False):
                continue

            try:
                ptype = shape.placeholder_format.type
            except Exception:
                continue

            if ptype not in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                continue

            tf = shape.text_frame
            tf.word_wrap = False
            tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

            had_break = False
            for para in tf.paragraphs:
                for run in para.runs:
                    if run.text and "\n" in run.text:
                        run.text = run.text.replace("\n", " ")
                        had_break = True
                        replaced_line_breaks += 1

            if "\n" in (shape.text or ""):
                shape.text = shape.text.replace("\n", " ")
                had_break = True

            if had_break:
                touched_titles += 1

    prs.save(output_path)
    return {
        "titles_touched": touched_titles,
        "title_line_breaks_replaced": replaced_line_breaks,
    }


def run_ppt_quality_checks(source_path: str, output_path: str, target_lang: str) -> PPTQAReport:
    src = Presentation(source_path)
    out = Presentation(output_path)

    out_rows = _iter_paragraph_texts(out)
    src_rows = _iter_paragraph_texts(src)
    source_lang = _detect_source_language(src_rows)

    issues: List[str] = []
    stats: Dict[str, int] = {
        "slide_count_mismatch": 0,
        "shape_count_mismatch": 0,
        "residual_kana": 0,
        "double_pattern_hits": 0,
        "table_prefix_hits": 0,
        "unchanged_segments": 0,
        "cjk_source_segments": 0,
        "output_segments": len(out_rows),
        "source_residual_segments": 0,
        "title_newline_hits": 0,
        "color_mismatch": 0,
    }

    if len(src.slides) != len(out.slides):
        stats["slide_count_mismatch"] = abs(len(src.slides) - len(out.slides))
        issues.append(f"页数不一致: source={len(src.slides)} output={len(out.slides)}")

    for i in range(min(len(src.slides), len(out.slides))):
        src_shape_count = len(list(_iter_shapes_recursive(src.slides[i].shapes)))
        out_shape_count = len(list(_iter_shapes_recursive(out.slides[i].shapes)))
        if src_shape_count != out_shape_count:
            stats["shape_count_mismatch"] += 1
            issues.append(f"[P{i+1}] 形状数不一致: source={src_shape_count} output={out_shape_count}")

    for kind, text, slide_idx in out_rows:
        for pattern, label in DOUBLE_PATTERNS:
            if pattern.search(text):
                stats["double_pattern_hits"] += 1
                issues.append(f"[P{slide_idx}] {label}: {text[:80]}")

        if target_lang == "Chinese" and _JP_KANA_RE.search(text):
            stats["residual_kana"] += 1
            issues.append(f"[P{slide_idx}] 残留日文假名: {text[:80]}")

        if source_lang == "English" and target_lang != "English" and _LATIN_RE.search(text):
            stats["source_residual_segments"] += 1
        elif source_lang == "Japanese" and target_lang != "Japanese" and _JP_KANA_RE.search(text):
            stats["source_residual_segments"] += 1
        elif source_lang == "Chinese" and target_lang == "English" and _HAN_RE.search(text):
            stats["source_residual_segments"] += 1

        if kind == "cell" and TABLE_BULLET_PREFIX.match(text):
            stats["table_prefix_hits"] += 1
            issues.append(f"[P{slide_idx}] 表格多余前缀: {text[:80]}")

    pair_count = min(len(src_rows), len(out_rows))
    for i in range(pair_count):
        _, src_text, _ = src_rows[i]
        _, out_text, _ = out_rows[i]
        if _CJK_RE.search(src_text):
            stats["cjk_source_segments"] += 1
            if src_text.strip() == out_text.strip():
                stats["unchanged_segments"] += 1

    for slide_idx, slide in enumerate(out.slides, 1):
        for shape in _iter_shapes_recursive(slide.shapes):
            if not getattr(shape, "has_text_frame", False) or not getattr(shape, "is_placeholder", False):
                continue
            try:
                ptype = shape.placeholder_format.type
            except Exception:
                continue
            if ptype in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                if "\n" in (shape.text or ""):
                    stats["title_newline_hits"] += 1
                    issues.append(f"[P{slide_idx}] 大标题存在换行: {(shape.text or '')[:80]}")

    src_colors = _collect_first_run_colors(src)
    out_colors = _collect_first_run_colors(out)
    
    if len(src_colors) != len(out_colors):
        stats["color_mismatch"] += abs(len(src_colors) - len(out_colors))
        issues.append(f"段落数变化导致颜色无法完整对比: 源={len(src_colors)} vs 输出={len(out_colors)}")
    
    c_pair = min(len(src_colors), len(out_colors))
    for i in range(c_pair):
        if src_colors[i] != out_colors[i]:
            stats["color_mismatch"] += 1

    unchanged_ratio = 0.0
    if stats["cjk_source_segments"] > 0:
        unchanged_ratio = stats["unchanged_segments"] / stats["cjk_source_segments"]
        if unchanged_ratio > UNCHANGED_RATIO_THRESHOLD:
            issues.append(
                f"疑似欠翻译比例偏高: {stats['unchanged_segments']}/{stats['cjk_source_segments']}"
                f" ({unchanged_ratio:.1%})"
            )

    source_residual_ratio = 0.0
    if source_lang != "Unknown" and source_lang != target_lang and stats["output_segments"] > 0:
        source_residual_ratio = stats["source_residual_segments"] / stats["output_segments"]
        if source_residual_ratio > SOURCE_RESIDUAL_RATIO_THRESHOLD:
            issues.append(
                f"源语言残留偏高({source_lang}): {stats['source_residual_segments']}/{stats['output_segments']} ({source_residual_ratio:.1%})"
            )

    translation_accuracy_ok = (
        source_residual_ratio <= SOURCE_RESIDUAL_RATIO_THRESHOLD
        and stats["double_pattern_hits"] == 0
        and unchanged_ratio <= UNCHANGED_RATIO_THRESHOLD
    )
    format_layout_ok = (
        stats["table_prefix_hits"] == 0
        and stats["slide_count_mismatch"] == 0
        and stats["shape_count_mismatch"] == 0
    )
    text_color_ok = stats["color_mismatch"] == 0
    title_single_line_ok = stats["title_newline_hits"] == 0

    if stats["color_mismatch"] > 0:
        issues.append(f"文本颜色疑似变化: {stats['color_mismatch']} 处")

    passed = translation_accuracy_ok and format_layout_ok and text_color_ok and title_single_line_ok

    if passed:
        summary = "QA通过：翻译准确性代理检查、排版代理检查、颜色一致性、标题单行检查均通过。"
    else:
        summary = "QA未完全通过：请依据问题列表进行人工复核。"

    return PPTQAReport(
        passed=passed,
        translation_accuracy_ok=translation_accuracy_ok,
        format_layout_ok=format_layout_ok,
        text_color_ok=text_color_ok,
        title_single_line_ok=title_single_line_ok,
        summary=summary,
        issues=issues,
        stats=stats,
        source_lang_detected=source_lang,
        target_lang=target_lang,
    )
