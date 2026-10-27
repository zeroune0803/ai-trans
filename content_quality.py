import re
from typing import Dict, List, Tuple

from docx import Document
from openpyxl import load_workbook

_CJK_RE = re.compile(r"[\u3000-\u9fff\uac00-\ud7af\uf900-\ufaff]")
_HAN_RE = re.compile(r"[\u4e00-\u9fff]")
_JP_KANA_RE = re.compile(r"[\u3041-\u3096\u30a1-\u30fa\u30fd-\u30ff]")
_LATIN_RE = re.compile(r"[A-Za-z]")

UNCHANGED_RATIO_THRESHOLD = 0.12
EN_CJK_RATIO_THRESHOLD = 0.08
JP_LATIN_RATIO_THRESHOLD = 0.08
SOURCE_RESIDUAL_RATIO_THRESHOLD = 0.08


def _detect_source_language(texts: List[str]) -> str:
    han_hits = 0
    kana_hits = 0
    latin_hits = 0

    for text in texts:
        if _HAN_RE.search(text):
            han_hits += 1
        if _JP_KANA_RE.search(text):
            kana_hits += 1
        if _LATIN_RE.search(text):
            latin_hits += 1

    if han_hits > 0:
        return "Chinese"
    if kana_hits > 0:
        return "Japanese"
    if latin_hits > 0:
        return "English"
    return "Unknown"


def _iter_docx_texts(path: str) -> List[str]:
    doc = Document(path)
    rows: List[str] = []

    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            rows.append(t)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = (cell.text or "").strip()
                if t:
                    rows.append(t)

    return rows


def _iter_xlsx_texts(path: str) -> List[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows: List[str] = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for val in row:
                    if val is None:
                        continue
                    t = str(val).strip()
                    if t:
                        rows.append(t)
    finally:
        wb.close()

    return rows


def _calc_content_quality(source_texts: List[str], output_texts: List[str], target_lang: str, file_type: str) -> Dict:
    issues: List[str] = []
    source_lang = _detect_source_language(source_texts)
    stats = {
        "source_segments": len(source_texts),
        "output_segments": len(output_texts),
        "cjk_source_segments": 0,
        "unchanged_segments": 0,
        "residual_kana": 0,
        "english_cjk_segments": 0,
        "japanese_latin_segments": 0,
        "source_residual_segments": 0,
    }

    pair_count = min(len(source_texts), len(output_texts))
    for i in range(pair_count):
        src_text = source_texts[i]
        out_text = output_texts[i]

        if _CJK_RE.search(src_text):
            stats["cjk_source_segments"] += 1
            if src_text.strip() == out_text.strip():
                stats["unchanged_segments"] += 1

        if target_lang == "Chinese" and _JP_KANA_RE.search(out_text):
            stats["residual_kana"] += 1

        if target_lang == "English" and _CJK_RE.search(out_text):
            stats["english_cjk_segments"] += 1

        if target_lang == "Japanese" and _LATIN_RE.search(out_text):
            stats["japanese_latin_segments"] += 1

        if source_lang == "English" and target_lang != "English" and _LATIN_RE.search(out_text):
            stats["source_residual_segments"] += 1
        elif source_lang == "Japanese" and target_lang != "Japanese" and _JP_KANA_RE.search(out_text):
            stats["source_residual_segments"] += 1
        elif source_lang == "Chinese" and target_lang == "English" and _HAN_RE.search(out_text):
            stats["source_residual_segments"] += 1

    unchanged_ratio = 0.0
    if stats["cjk_source_segments"] > 0:
        unchanged_ratio = stats["unchanged_segments"] / stats["cjk_source_segments"]
        if unchanged_ratio > UNCHANGED_RATIO_THRESHOLD:
            issues.append(
                f"疑似欠翻译比例偏高: {stats['unchanged_segments']}/{stats['cjk_source_segments']} ({unchanged_ratio:.1%})"
            )

    if target_lang == "Chinese" and stats["residual_kana"] > 0:
        issues.append(f"检测到日文假名残留: {stats['residual_kana']} 处")

    english_cjk_ratio = 0.0
    if target_lang == "English" and stats["output_segments"] > 0:
        english_cjk_ratio = stats["english_cjk_segments"] / stats["output_segments"]
        if english_cjk_ratio > EN_CJK_RATIO_THRESHOLD:
            issues.append(
                f"目标为 English 但输出仍含较多 CJK 文本: {stats['english_cjk_segments']}/{stats['output_segments']} ({english_cjk_ratio:.1%})"
            )

    japanese_latin_ratio = 0.0
    if target_lang == "Japanese" and stats["output_segments"] > 0:
        japanese_latin_ratio = stats["japanese_latin_segments"] / stats["output_segments"]
        if japanese_latin_ratio > JP_LATIN_RATIO_THRESHOLD:
            issues.append(
                f"目标为 Japanese 但输出仍含较多拉丁字母: {stats['japanese_latin_segments']}/{stats['output_segments']} ({japanese_latin_ratio:.1%})"
            )

    source_residual_ratio = 0.0
    if source_lang != "Unknown" and source_lang != target_lang and stats["output_segments"] > 0:
        source_residual_ratio = stats["source_residual_segments"] / stats["output_segments"]
        if source_residual_ratio > SOURCE_RESIDUAL_RATIO_THRESHOLD:
            issues.append(
                f"源语言残留偏高({source_lang}): {stats['source_residual_segments']}/{stats['output_segments']} ({source_residual_ratio:.1%})"
            )

    translation_accuracy_ok = not issues
    passed = translation_accuracy_ok

    if passed:
        summary = f"{file_type} QA通过：内容翻译一致性检查通过。"
    else:
        summary = f"{file_type} QA未完全通过：请依据问题列表进行人工复核。"

    return {
        "passed": passed,
        "summary": summary,
        "issues": issues,
        "stats": stats,
        "source_lang_detected": source_lang,
        "target_lang": target_lang,
        "translation_accuracy_ok": translation_accuracy_ok,
        "checks_skipped": ["layout", "color", "title"],
    }


def run_docx_quality_checks(source_path: str, output_path: str, target_lang: str) -> Dict:
    src_rows = _iter_docx_texts(source_path)
    out_rows = _iter_docx_texts(output_path)
    return _calc_content_quality(src_rows, out_rows, target_lang, file_type="DOCX")


def run_xlsx_quality_checks(source_path: str, output_path: str, target_lang: str) -> Dict:
    src_rows = _iter_xlsx_texts(source_path)
    out_rows = _iter_xlsx_texts(output_path)
    return _calc_content_quality(src_rows, out_rows, target_lang, file_type="XLSX")
