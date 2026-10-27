import argparse
import os
import sys
import json
import csv
import glob
from collections import defaultdict

from openpyxl import load_workbook

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from translate_ppt_batch import translate_items_in_parallel


def norm(v):
    return str(v).strip() if v is not None else ""


def detect_language_columns(header_row):
    cols = {}
    for idx, cell in enumerate(header_row, start=1):
        h = norm(cell).lower()
        if not h:
            continue
        if any(k in h for k in ["中", "中文", "chinese", "cn", "source", "原文", "源文"]):
            cols.setdefault("zh", idx)
        if any(k in h for k in ["日", "日文", "日语", "japanese", "jp", "ja", "目标文", "译文"]):
            cols.setdefault("ja", idx)
        if any(k in h for k in ["英", "英文", "英语", "english", "en"]):
            cols.setdefault("en", idx)
    return cols


def find_header_row(ws, max_scan=10):
    for r in range(1, min(max_scan, ws.max_row) + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        cols = detect_language_columns(values)
        if len(cols) >= 2:
            return r, cols
    return None, {}


def build_pair_maps(ws, header_row, lang_cols):
    pair_maps = defaultdict(dict)
    for r in range(header_row + 1, ws.max_row + 1):
        vals = {lang: norm(ws.cell(r, c).value) for lang, c in lang_cols.items()}
        for src_lang, src_text in vals.items():
            if not src_text:
                continue
            for dst_lang, dst_text in vals.items():
                if src_lang == dst_lang or not dst_text:
                    continue
                if src_text not in pair_maps[(src_lang, dst_lang)]:
                    pair_maps[(src_lang, dst_lang)][src_text] = dst_text
    return pair_maps


def load_external_maps(base_dir):
    pair_maps = defaultdict(dict)
    for path in sorted(glob.glob(os.path.join(base_dir, "glossary_ja*.json"))):
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, dict):
                for k, v in data.items():
                    src = norm(k)
                    dst = norm(v)
                    if src and dst:
                        pair_maps[("zh", "ja")][src] = dst
                        pair_maps[("ja", "zh")][dst] = src
        except Exception:
            pass

    for path in sorted(glob.glob(os.path.join(base_dir, "glossary_ja*.csv"))):
        try:
            with open(path, "r", encoding="utf-8-sig") as f:
                reader = csv.reader(f)
                for row in reader:
                    if len(row) < 2:
                        continue
                    src = norm(row[0])
                    dst = norm(row[1])
                    if src and dst:
                        pair_maps[("zh", "ja")][src] = dst
                        pair_maps[("ja", "zh")][dst] = src
        except Exception:
            pass

    return pair_maps


def build_global_pair_maps(wb):
    global_maps = defaultdict(dict)
    for ws in wb.worksheets:
        header_row, lang_cols = find_header_row(ws)
        if not header_row or len(lang_cols) < 2:
            continue
        local_maps = build_pair_maps(ws, header_row, lang_cols)
        for key, mp in local_maps.items():
            for src, dst in mp.items():
                if src and dst and src not in global_maps[key]:
                    global_maps[key][src] = dst
    return global_maps


def fill_sheet(ws, pair_maps):
    header_row, lang_cols = find_header_row(ws)
    if not header_row or len(lang_cols) < 2:
        return 0

    filled = 0

    for r in range(header_row + 1, ws.max_row + 1):
        vals = {lang: norm(ws.cell(r, c).value) for lang, c in lang_cols.items()}

        for dst_lang, dst_col in lang_cols.items():
            if vals.get(dst_lang):
                continue

            candidate = ""
            for src_lang, src_text in vals.items():
                if src_lang == dst_lang or not src_text:
                    continue
                candidate = pair_maps.get((src_lang, dst_lang), {}).get(src_text, "")
                if candidate:
                    break

            if candidate:
                ws.cell(r, dst_col).value = candidate
                vals[dst_lang] = candidate
                filled += 1

    return filled


def count_remaining(ws):
    header_row, lang_cols = find_header_row(ws)
    if not header_row or "ja" not in lang_cols:
        return 0

    ja_col = lang_cols["ja"]
    count = 0
    for r in range(header_row + 1, ws.max_row + 1):
        ja_val = norm(ws.cell(r, ja_col).value)
        if ja_val:
            continue
        zh_val = norm(ws.cell(r, lang_cols.get("zh", 0)).value) if "zh" in lang_cols else ""
        en_val = norm(ws.cell(r, lang_cols.get("en", 0)).value) if "en" in lang_cols else ""
        if zh_val or en_val:
            count += 1
    return count


def collect_missing_ja_cells(wb):
    missing = []
    item_id = 1
    for ws in wb.worksheets:
        header_row, lang_cols = find_header_row(ws)
        if not header_row or "ja" not in lang_cols:
            continue

        ja_col = lang_cols["ja"]
        for r in range(header_row + 1, ws.max_row + 1):
            ja_val = norm(ws.cell(r, ja_col).value)
            if ja_val:
                continue

            zh_val = norm(ws.cell(r, lang_cols.get("zh", 0)).value) if "zh" in lang_cols else ""
            en_val = norm(ws.cell(r, lang_cols.get("en", 0)).value) if "en" in lang_cols else ""
            src = zh_val or en_val
            if not src:
                continue

            missing.append({
                "id": item_id,
                "text": src,
                "ws": ws,
                "row": r,
                "col": ja_col,
            })
            item_id += 1
    return missing


def fill_missing_with_ai(wb, glossary_zh_ja):
    missing = collect_missing_ja_cells(wb)
    if not missing:
        return 0

    text_items = [{"id": m["id"], "text": m["text"]} for m in missing]
    translated, _, _ = translate_items_in_parallel(
        text_items=text_items,
        glossary=glossary_zh_ja,
        target_lang="Japanese",
        speed_mode="balanced",
        max_retries=2,
        enable_quality_gate=True,
        batch_size=20,
        progress_callback=None,
        progress_stage="术语补全翻译中...",
        log_prefix="[GLOSSARY]",
    )

    filled = 0
    by_id = {m["id"]: m for m in missing}
    for k, v in translated.items():
        m = by_id.get(k)
        dst = norm(v)
        if m and dst:
            m["ws"].cell(m["row"], m["col"]).value = dst
            filled += 1
    return filled


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input_path")
    parser.add_argument("--output", default="")
    parser.add_argument("--ai", action="store_true")
    args = parser.parse_args()

    input_path = args.input_path
    if not os.path.exists(input_path):
        raise FileNotFoundError(input_path)

    wb = load_workbook(input_path)
    pair_maps = build_global_pair_maps(wb)
    external_maps = load_external_maps(PROJECT_ROOT)
    for key, mp in external_maps.items():
        pair_maps[key].update(mp)

    total_filled = 0
    for ws in wb.worksheets:
        total_filled += fill_sheet(ws, pair_maps)

    ai_filled = 0
    if args.ai:
        ai_filled = fill_missing_with_ai(wb, pair_maps.get(("zh", "ja"), {}))

    remaining_ja = 0
    for ws in wb.worksheets:
        remaining_ja += count_remaining(ws)

    output_path = args.output
    if not output_path:
        base, ext = os.path.splitext(input_path)
        output_path = f"{base}_已补全{ext}"

    wb.save(output_path)
    print(f"FILLED_BY_MAP={total_filled}")
    print(f"FILLED_BY_AI={ai_filled}")
    print(f"FILLED_TOTAL={total_filled + ai_filled}")
    print(f"REMAINING_JA_EMPTY_WITH_SOURCE={remaining_ja}")
    print(f"OUTPUT={output_path}")


if __name__ == "__main__":
    main()
