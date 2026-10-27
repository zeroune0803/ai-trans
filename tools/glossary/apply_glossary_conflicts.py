import csv
import os
import sys

from openpyxl import load_workbook


def norm(v):
    return str(v).strip() if v is not None else ""


def detect_language_columns(header_row):
    cols = {}
    for idx, cell in enumerate(header_row, start=1):
        h = norm(cell).lower()
        if not h:
            continue
        if any(k in h for k in ["英", "英文", "英语", "english", "en"]):
            cols.setdefault("en", idx)
        if any(k in h for k in ["中", "中文", "chinese", "cn", "source", "原文", "源文"]):
            cols.setdefault("zh", idx)
        if any(k in h for k in ["日", "日文", "日语", "japanese", "jp", "ja", "目标文", "译文"]):
            cols.setdefault("ja", idx)
    return cols


def find_header_row(ws, max_scan=10):
    for r in range(1, min(max_scan, ws.max_row) + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        cols = detect_language_columns(values)
        if "zh" in cols:
            return r, cols
    return None, {}


def load_conflict_map(csv_path, target_lang):
    mapping = {}
    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        target_col = "日文译法" if target_lang == "ja" else "英文译法"
        for row in reader:
            zh = norm(row.get("中文术语"))
            target_text = norm(row.get(target_col))
            if not zh or not target_text:
                continue
            mapping[zh] = target_text
    return mapping


def apply_mapping(wb, mapping, target_lang):
    changed = 0
    touched_terms = set()

    for ws in wb.worksheets:
        header_row, cols = find_header_row(ws)
        if not header_row:
            continue
        if target_lang not in cols:
            continue

        zh_col = cols["zh"]
        target_col = cols[target_lang]

        for r in range(header_row + 1, ws.max_row + 1):
            zh = norm(ws.cell(r, zh_col).value)
            if not zh or zh not in mapping:
                continue

            target_text = mapping[zh]
            current_text = norm(ws.cell(r, target_col).value)
            if current_text != target_text:
                ws.cell(r, target_col).value = target_text
                changed += 1
            touched_terms.add(zh)

    return changed, len(touched_terms)


def main():
    if len(sys.argv) < 3:
        print("Usage: python apply_glossary_conflicts.py <xlsx_path> <conflict_csv_path> [output_xlsx] [target_lang]")
        sys.exit(1)

    xlsx_path = sys.argv[1]
    conflict_csv = sys.argv[2]
    output_xlsx = sys.argv[3] if len(sys.argv) >= 4 else ""
    target_lang = (sys.argv[4] if len(sys.argv) >= 5 else "ja").strip().lower()
    if target_lang not in {"ja", "en"}:
        raise ValueError("target_lang must be 'ja' or 'en'")

    if not os.path.exists(xlsx_path):
        raise FileNotFoundError(xlsx_path)
    if not os.path.exists(conflict_csv):
        raise FileNotFoundError(conflict_csv)

    mapping = load_conflict_map(conflict_csv, target_lang)
    wb = load_workbook(xlsx_path)
    changed, touched = apply_mapping(wb, mapping, target_lang)

    if not output_xlsx:
        base, ext = os.path.splitext(xlsx_path)
        output_xlsx = f"{base}_已统一{ext}"

    wb.save(output_xlsx)
    print(f"MAPPING_TERMS={len(mapping)}")
    print(f"TOUCHED_TERMS={touched}")
    print(f"CHANGED_CELLS={changed}")
    print(f"OUTPUT={output_xlsx}")


if __name__ == "__main__":
    main()
