import csv
import os
import sys
from collections import defaultdict

from openpyxl import load_workbook


def norm(v):
    return str(v).strip() if v is not None else ""


def detect_language_columns(header_row):
    cols = {}
    for idx, cell in enumerate(header_row, start=1):
        h = norm(cell).lower()
        if not h:
            continue
        if any(k in h for k in ["英", "英文", "英语", "english", "en"]):
            cols.setdefault("en", idx)
        if any(k in h for k in ["中", "中文", "chinese", "cn", "source", "原文", "源文"]):
            cols.setdefault("zh", idx)
        if any(k in h for k in ["日", "日文", "日语", "japanese", "jp", "ja", "目标文", "译文"]):
            cols.setdefault("ja", idx)
    return cols


def find_header_row(ws, max_scan=10):
    for r in range(1, min(max_scan, ws.max_row) + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        cols = detect_language_columns(values)
        if "zh" in cols:
            return r, cols
    return None, {}


def main():
    if len(sys.argv) < 2:
        print("Usage: python check_glossary_consistency.py <xlsx_path> [output_csv] [target_lang]")
        sys.exit(1)

    xlsx_path = sys.argv[1]
    out_csv = sys.argv[2] if len(sys.argv) > 2 else ""
    target_lang = (sys.argv[3] if len(sys.argv) > 3 else "ja").strip().lower()
    if target_lang not in {"ja", "en"}:
        raise ValueError("target_lang must be 'ja' or 'en'")

    wb = load_workbook(xlsx_path, data_only=True)

    zh_to_target = defaultdict(set)
    details = defaultdict(list)

    for ws in wb.worksheets:
        header_row, cols = find_header_row(ws)
        if not header_row:
            continue
        if target_lang not in cols:
            continue

        zh_col = cols["zh"]
        target_col = cols[target_lang]
        for r in range(header_row + 1, ws.max_row + 1):
            zh = norm(ws.cell(r, zh_col).value)
            target_text = norm(ws.cell(r, target_col).value)
            if not zh or not target_text:
                continue
            zh_to_target[zh].add(target_text)
            details[zh].append((target_text, ws.title, r))

    conflicts = []
    for zh, target_set in zh_to_target.items():
        if len(target_set) > 1:
            rows = details[zh]
            conflicts.append((zh, sorted(target_set), rows))

    conflicts.sort(key=lambda x: len(x[1]), reverse=True)

    print(f"TARGET_LANG={target_lang}")
    print(f"TOTAL_ZH_TERMS={len(zh_to_target)}")
    print(f"CONFLICT_TERMS={len(conflicts)}")

    for i, (zh, target_list, rows) in enumerate(conflicts[:20], start=1):
        print(f"[{i}] {zh}")
        print(f"    TARGET_VARIANTS={len(target_list)} -> {' | '.join(target_list)}")

    if not out_csv:
        base, _ = os.path.splitext(xlsx_path)
        out_csv = base + f"_术语一致性冲突_{target_lang}.csv"

    with open(out_csv, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        target_col_name = "日文译法" if target_lang == "ja" else "英文译法"
        writer.writerow(["中文术语", target_col_name, "变体数量", "出现位置"])
        for zh, target_list, rows in conflicts:
            # aggregate positions by variant
            by_target = defaultdict(list)
            for target_text, sheet, row in rows:
                by_target[target_text].append(f"{sheet}!R{row}")
            for target_text in target_list:
                pos = "; ".join(by_target[target_text][:20])
                writer.writerow([zh, target_text, len(target_list), pos])

    print(f"OUTPUT={out_csv}")


if __name__ == "__main__":
    main()
