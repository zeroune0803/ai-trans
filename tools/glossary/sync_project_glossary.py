import json
import os
import sys
from collections import defaultdict

from openpyxl import load_workbook


def norm(v):
    if v is None:
        return ""
    t = str(v)
    for ch in ("\ufeff", "\u200b", "\u200c", "\u200d", "\u2060"):
        t = t.replace(ch, "")
    t = t.replace("\xa0", " ")
    return t.strip()


def detect_language_columns(header_row):
    cols = {}
    for idx, cell in enumerate(header_row, start=1):
        h = norm(cell).lower()
        if not h:
            continue
        if any(k in h for k in ["英", "英文", "英语", "english", "en"]):
            cols.setdefault("en", idx)
        if any(k in h for k in ["中", "中文", "chinese", "cn", "source", "原文", "源文"]):
            cols.setdefault("zh", idx)
        if any(k in h for k in ["日", "日文", "日语", "japanese", "jp", "ja", "目标文", "译文"]):
            cols.setdefault("ja", idx)
    return cols


def find_header_row(ws, max_scan=10):
    for r in range(1, min(max_scan, ws.max_row) + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        cols = detect_language_columns(values)
        if "zh" in cols and ("ja" in cols or "en" in cols):
            return r, cols
    return None, {}


def collect_maps(xlsx_path):
    wb = load_workbook(xlsx_path, data_only=True)

    zh_ja = {}
    zh_en = {}
    conflicts_ja = defaultdict(set)
    conflicts_en = defaultdict(set)

    for ws in wb.worksheets:
        header_row, cols = find_header_row(ws)
        if not header_row:
            continue

        zh_col = cols.get("zh")
        ja_col = cols.get("ja")
        en_col = cols.get("en")

        for r in range(header_row + 1, ws.max_row + 1):
            zh = norm(ws.cell(r, zh_col).value) if zh_col else ""
            if not zh:
                continue

            ja = norm(ws.cell(r, ja_col).value) if ja_col else ""
            en = norm(ws.cell(r, en_col).value) if en_col else ""

            if ja:
                conflicts_ja[zh].add(ja)
                zh_ja[zh] = ja
            if en:
                conflicts_en[zh].add(en)
                zh_en[zh] = en

    conflict_terms_ja = sum(1 for v in conflicts_ja.values() if len(v) > 1)
    conflict_terms_en = sum(1 for v in conflicts_en.values() if len(v) > 1)

    return zh_ja, zh_en, conflict_terms_ja, conflict_terms_en


def write_json(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)


def main():
    if len(sys.argv) < 2:
        print("Usage: python sync_project_glossary.py <xlsx_path>")
        sys.exit(1)

    xlsx_path = sys.argv[1]
    if not os.path.exists(xlsx_path):
        raise FileNotFoundError(xlsx_path)

    project_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
    ja_out = os.path.join(project_dir, "glossary_ja_from_table.json")
    en_out = os.path.join(project_dir, "glossary_en_from_table.json")

    zh_ja, zh_en, c_ja, c_en = collect_maps(xlsx_path)

    write_json(ja_out, zh_ja)
    write_json(en_out, zh_en)

    print(f"SOURCE={xlsx_path}")
    print(f"JA_TERMS={len(zh_ja)}")
    print(f"EN_TERMS={len(zh_en)}")
    print(f"JA_CONFLICT_TERMS_IN_SOURCE={c_ja}")
    print(f"EN_CONFLICT_TERMS_IN_SOURCE={c_en}")
    print(f"JA_OUT={ja_out}")
    print(f"EN_OUT={en_out}")


if __name__ == "__main__":
    main()
