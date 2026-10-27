import sys
from openpyxl import load_workbook

path = sys.argv[1]
wb = load_workbook(path, data_only=True)
print("FILE:", path)
for ws in wb.worksheets:
    print("SHEET:", ws.title, "rows=", ws.max_row, "cols=", ws.max_column)
    max_c = min(ws.max_column, 12)
    for r in range(1, min(ws.max_row, 8) + 1):
        vals = [ws.cell(r, c).value for c in range(1, max_c + 1)]
        print(r, vals)
    print("-" * 80)
