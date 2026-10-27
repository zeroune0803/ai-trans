import os
from copy import copy
import math

import openpyxl
from openpyxl.styles import Font
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from translate_ppt_batch import (
    generate_glossary,
    translate_items_in_parallel,
)
from rules import get_target_font_name, resolve_speed_options, sanitize_glossary_for_target


def _set_cell_font(cell, target_lang):
    font_name = get_target_font_name(target_lang, content_type="presentation")
    if not font_name:
        return

    if cell.font:
        new_font = copy(cell.font)
        new_font.name = font_name
        cell.font = new_font
    else:
        cell.font = Font(name=font_name)


def _approx_cell_capacity(cell):
    """Approximate characters that can fit one line for current column width."""
    ws = cell.parent
    col_letter = get_column_letter(cell.column)
    col_dim = ws.column_dimensions.get(col_letter)
    col_w = col_dim.width if (col_dim and col_dim.width) else 8.43
    # Excel width unit to characters is roughly near 1:1 for Latin; CJK less dense.
    return max(6, int(col_w * 1.15))


def _repair_cell_layout(cell, src_text, dst_text):
    src_len = max(1, len((src_text or "").strip()))
    dst_len = len((dst_text or "").strip())
    if dst_len <= src_len:
        return False

    cap = _approx_cell_capacity(cell)
    risky = dst_len > int(cap * 1.25) or dst_len > int(src_len * 1.5)
    if not risky:
        return False

    try:
        old_align = copy(cell.alignment) if cell.alignment else Alignment()
    except Exception:
        old_align = Alignment()

    try:
        cell.alignment = Alignment(
            horizontal=old_align.horizontal,
            vertical=old_align.vertical,
            text_rotation=old_align.text_rotation,
            wrap_text=True,
            shrink_to_fit=True,
            indent=old_align.indent,
            relative_indent=old_align.relative_indent,
            justify_last_line=old_align.justify_last_line,
            reading_order=old_align.reading_order,
            merge_cell=old_align.merge_cell,
        )
    except Exception:
        pass

    try:
        if cell.font:
            f = copy(cell.font)
            base_size = float(f.sz) if f.sz else 11.0
            ratio = max(1.0, dst_len / max(cap, 1))
            new_size = max(8.0, min(base_size, base_size / (ratio ** 0.45)))
            f.sz = new_size
            cell.font = f
    except Exception:
        pass

    return True


def _maybe_expand_column_width(sheet, col_idx, text):
    """Gentle width grow to reduce severe overflow while limiting layout drift."""
    col_letter = get_column_letter(col_idx)
    dim = sheet.column_dimensions[col_letter]
    current = dim.width if dim.width else 8.43
    desired = min(32.0, max(current, len((text or "")) * 0.62))
    if desired > current + 0.5:
        dim.width = desired


def translate_excel_workbook(
    input_file,
    output_file,
    target_lang="Chinese",
    speed_mode="balanced",
    progress_callback=None,
):
    print(f"正在加载 Excel 文件: {input_file}")
    try:
        wb = openpyxl.load_workbook(input_file)
    except Exception as e:
        print(f"加载工作簿失败: {e}")
        return

    text_items = []
    pointers = {}
    idx_ref = 0

    total_sheets = len(wb.worksheets)
    for sheet_idx, sheet in enumerate(wb.worksheets, 1):
        print(f"[XLSX] 扫描 Sheet {sheet_idx}/{total_sheets}: {sheet.title}")
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                src = cell.value.strip()
                if not src or src.startswith("="):
                    continue

                text_items.append({"id": idx_ref, "text": src, "sheet_idx": sheet_idx})
                pointers[idx_ref] = cell
                idx_ref += 1

    if not text_items:
        print("[XLSX] 未发现可翻译文本，直接保存")
        wb.save(output_file)
        return

    mode_opts = resolve_speed_options(speed_mode)
    batch_size = mode_opts["batch_size"]
    max_retries = mode_opts["max_retries"]
    enable_quality_gate = mode_opts["enable_quality_gate"]

    source_texts = [x["text"] for x in text_items]
    glossary = generate_glossary(source_texts, target_lang)
    glossary = sanitize_glossary_for_target(glossary, target_lang, source_texts)

    translated, last_model_used, used_models = translate_items_in_parallel(
        text_items,
        glossary,
        target_lang,
        speed_mode=speed_mode,
        max_retries=max_retries,
        enable_quality_gate=enable_quality_gate,
        batch_size=batch_size,
        progress_callback=progress_callback,
        progress_stage="调用模型并行翻译中...",
        log_prefix="[XLSX]",
    )
    total_batches = (len(text_items) + batch_size - 1) // batch_size
    if progress_callback and last_model_used:
        progress_callback("调用模型并行翻译中...", total_batches, total_batches, model=last_model_used)
    if used_models:
        print(f"[XLSX] 本次任务命中模型: {', '.join(used_models)}")

    last_sheet_idx = None
    repaired = 0
    for item in text_items:
        sheet_idx = item["sheet_idx"]
        if progress_callback and sheet_idx != last_sheet_idx:
            progress_callback("工作表回填阶段", sheet_idx, total_sheets)
            last_sheet_idx = sheet_idx

        cell = pointers[item["id"]]
        cell.value = translated.get(item["id"], item["text"])
        _set_cell_font(cell, target_lang)
        if _repair_cell_layout(cell, item["text"], cell.value):
            repaired += 1
        _maybe_expand_column_width(cell.parent, cell.column, cell.value)

    if repaired > 0:
        print(f"[XLSX] 已触发自动排版修复 {repaired} 处（列宽/换行/缩放）。")

    print(f"保存翻译完毕的新文件到: {output_file}")
    wb.save(output_file)
    print("Excel翻译任务完成！")


if __name__ == "__main__":
    input_excel = "test.xlsx"
    output_excel = "test_translated.xlsx"
    if os.path.exists(input_excel):
        translate_excel_workbook(input_excel, output_excel, target_lang="Japanese")
    else:
        print(f"测试文件 {input_excel} 不存在，请放入测试文件。")
