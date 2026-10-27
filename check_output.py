from pptx import Presentation
import re
import sys

path = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\DMC86124042\Documents\WXWork\1688854748647373\Cache\File\2026-05\2-②_AIシフト最適化_Translated_Chinese.pptx"
prs = Presentation(path)

issues = []

double_patterns = [
    (r'AI\s*AI', 'AI重复'),
    (r'DX\s*DX', 'DX重复'),
    (r'APP\s*APP', 'APP重复'),
    (r'※\s*※', '※重复'),
    (r'(\d{4})\1', '4位数字重复(年份)'),
    (r'([①②③④⑤⑥⑦⑧⑨⑩❶❷❸❹❺❻❼❽❾❿])\1', '序号圆圈重复'),
]
bullet_prefix = re.compile(r'^[\s]*[•·●○◆▶→\-—]\s')
# 只检测真正的平假名/片假名字母，排除 ・(U+30FB 中点) 和 ー(U+30FC 长音符)
# 这两个符号在中文翻译文本里是合法分隔符，不属于"残留日文"
japanese_remaining = re.compile(r'[\u3041-\u3096\u30a1-\u30fa\u30fd-\u30ff]')

for slide_idx, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        items = []
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                items.append(('para', para.text.strip()))
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for para in cell.text_frame.paragraphs:
                        items.append(('cell', para.text.strip()))
        for kind, text in items:
            if not text:
                continue
            for pattern, label in double_patterns:
                if re.search(pattern, text):
                    issues.append(f"  [P{slide_idx}] {label}: {text[:80]}")
            if japanese_remaining.search(text):
                issues.append(f"  [P{slide_idx}] 残留日文假名: {text[:80]}")
            if kind == 'cell' and bullet_prefix.match(text):
                issues.append(f"  [P{slide_idx}] 表格多余前缀: {text[:80]}")

if issues:
    print(f"发现 {len(issues)} 个问题：")
    for i in issues:
        print(i)
else:
    print("✅ 自检通过：未发现文字重复、残留日文、表格多余前缀等问题。")
