import os
import json
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

source_file = r"C:\Users\DMC86124042\Downloads\ASPD L3运维未尽事宜沟通_202605_CN_Revised3.pptx"
target_file = r"C:\Users\DMC86124042\Downloads\ASPD L3运维未尽事宜沟通_202605_CN_Revised3_JP_v2.pptx"
dict_file = r"e:/Dev/translate_agent/trans_map_final.json"

def apply_translation():
    prs = Presentation(source_file)
    
    with open(dict_file, 'r', encoding='utf-8') as f:
        trans_map = json.load(f)

    # 递归提取所有包含文字的形状
    def process_shape(shape, is_title=False):
        if shape.has_text_frame:
            if is_title:
                shape.text_frame.word_wrap = False
            
            try:
                # 调整字体大小自动适应框宽
                shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            except:
                pass
                
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if text in trans_map:
                    translated_text = trans_map[text]
                    if len(paragraph.runs) > 0:
                        paragraph.runs[0].text = translated_text
                        for run in paragraph.runs[1:]:
                            run.text = ""

        if getattr(shape, 'has_table', False):
            for row in shape.table.rows:
                for cell in row.cells:
                    try:
                        cell.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                    except:
                        pass
                    for paragraph in cell.text_frame.paragraphs:
                        text = "".join(run.text for run in paragraph.runs).strip()
                        if text in trans_map:
                            translated_text = trans_map[text]
                            if len(paragraph.runs) > 0:
                                paragraph.runs[0].text = translated_text
                                for run in paragraph.runs[1:]:
                                    run.text = ""

        if shape.shape_type == 6: # Group shape
            for sub_shape in shape.shapes:
                # 递归处理
                process_shape(sub_shape)

    for slide in prs.slides:
        title_shape = slide.shapes.title if slide.shapes.title else None
        for shape in slide.shapes:
            process_shape(shape, is_title=(shape == title_shape))

    prs.save(target_file)
    print(f"Translation logic complete with format adjustments. Saved to {target_file}")

if __name__ == "__main__":
    apply_translation()