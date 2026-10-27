from pptx import Presentation
import re

src = Presentation(r'C:\Users\DMC86124042\Documents\WXWork\1688854748647373\Cache\File\2026-05\2-②_AIシフト最適化.pptx')
CJK_RE = re.compile(r'[\u3000-\u9fff]')
circles = re.compile(r'[❶❷❸❹❺❻①②③④⑤⑥]')

slide = list(src.slides)[2]  # index 2 = P3
print(f"=== 源文件 P3 所有含圆圈数字的 shape ===")
for shape in slide.shapes:
    if not shape.has_text_frame:
        continue
    for para in shape.text_frame.paragraphs:
        t = para.text
        if circles.search(t):
            print(f"\nshape_type={shape.shape_type} full_para_text={repr(t[:80])}")
            for i, r in enumerate(para.runs):
                if r.text:
                    print(f"  run[{i}]: {repr(r.text[:40])}  has_cjk={bool(CJK_RE.search(r.text))}")
