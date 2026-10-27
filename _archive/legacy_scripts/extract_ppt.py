import os
import json
from pptx import Presentation

source_file = r"C:\Users\DMC86124042\Downloads\ASPD L3运维未尽事宜沟通_202605_CN_Revised3.pptx"
extract_file = "e:/Dev/translate_agent/extracted_texts.json"

if not os.path.exists(source_file):
    print(f"File not found: {source_file}")
else:
    prs = Presentation(source_file)
    texts = []
    
    for slide_idx, slide in enumerate(prs.slides):
        for shape_idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                for para_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        texts.append({"original": text})
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text_frame:
                            for paragraph in cell.text_frame.paragraphs:
                                text = "".join(run.text for run in paragraph.runs).strip()
                                if text:
                                    texts.append({"original": text})
            if shape.shape_type == 6: # Grouped shape
                for sub_shape in shape.shapes:
                    if sub_shape.has_text_frame:
                        for paragraph in sub_shape.text_frame.paragraphs:
                            text = "".join(run.text for run in paragraph.runs).strip()
                            if text:
                                texts.append({"original": text})
                            
    with open(extract_file, 'w', encoding='utf-8') as f:
        json.dump(texts, f, ensure_ascii=False, indent=4)
        
    print(f"Extracted {len(texts)} text blocks to {extract_file}")
