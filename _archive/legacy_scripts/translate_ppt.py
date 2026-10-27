import os
from pptx import Presentation
from deep_translator import GoogleTranslator

# 源文件路径
source_file = r"C:\Users\DMC86124042\Downloads\ASPD L3运维未尽事宜沟通_202605_CN_Revised3.pptx"
# 目标文件路径
target_file = r"C:\Users\DMC86124042\Downloads\ASPD L3运维未尽事宜沟通_202605_CN_Revised3_JP.pptx"

def translate_ppt(input_path, output_path):
    print(f"正在读取文件: {input_path}")
    prs = Presentation(input_path)
    # 尝试使用不同的翻译器以防网络问题
    from deep_translator import GoogleTranslator, MyMemoryTranslator
    try:
        translator = GoogleTranslator(source='zh-CN', target='ja')
        # 测试一下网络
        translator.translate("测试")
        print("使用 GoogleTranslator")
    except:
        print("GoogleTranslator 失败，切换到 MyMemoryTranslator...")
        translator = MyMemoryTranslator(source='zh-CN', target='ja')

    total_shapes = 0
    # 遍历幻灯片
    for slide_idx, slide in enumerate(prs.slides):
        print(f"正在处理第 {slide_idx + 1} 页幻灯片...")
        # 遍历形状
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            
            # 遍历文本框中的段落
            for paragraph in shape.text_frame.paragraphs:
                # 遍历段落中的运行块，保留原有的格式和排版
                for run in paragraph.runs:
                    text = run.text.strip()
                    if text:
                        try:
                            # 翻译文本
                            translated_text = translator.translate(text)
                            run.text = translated_text
                            total_shapes += 1
                        except Exception as e:
                            print(f"翻译失败: '{text}' - 错误: {e}")

    # 保存翻译后的文件
    prs.save(output_path)
    print(f"\n翻译完成！共翻译了 {total_shapes} 个文本块。")
    print(f"文件已保存至: {output_path}")

if __name__ == "__main__":
    if os.path.exists(source_file):
        translate_ppt(source_file, target_file)
    else:
        print(f"找不到源文件: {source_file}")
