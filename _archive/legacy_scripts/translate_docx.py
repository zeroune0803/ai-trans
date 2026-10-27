import os
from docx import Document
from deep_translator import GoogleTranslator
import time

source_file = r"C:\Users\DMC86124042\Downloads\word《存量时代下的零售分化与数字化价值选择》.docx"
target_file = r"C:\Users\DMC86124042\Downloads\word《存量时代下的零售分化与数字化价值选择》_JP.docx"

translator = GoogleTranslator(source='zh-CN', target='ja')

# 载入长期记忆词典
memory_dict = {
    "双社长会": "ステコミ会議"
}

def translate_text(text):
    text_stripped = text.strip()
    if not text_stripped:
        return text
    
    # 应用长期记忆
    for k, v in memory_dict.items():
        text = text.replace(k, v)
        
    try:
        translated = translator.translate(text)
        return translated
    except Exception as e:
        print(f"Translation Error for '{text_stripped}': {e}")
        return text

def apply_translation():
    print(f"正在读取 Word 文档: {source_file}")
    doc = Document(source_file)
    total_translated = 0
    
    # 1. 翻译正文段落
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            original_text = paragraph.text
            translated_text = translate_text(original_text)
            
            if translated_text != original_text:
                # 尽量保留段落主要格式：将翻译文本写入第一个运行块（run），清空其他
                if len(paragraph.runs) > 0:
                    paragraph.runs[0].text = translated_text
                    for run in paragraph.runs[1:]:
                        run.text = ""
                    total_translated += 1

    # 2. 翻译表格内文字
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip():
                        original_text = paragraph.text
                        translated_text = translate_text(original_text)
                        
                        if translated_text != original_text:
                            if len(paragraph.runs) > 0:
                                paragraph.runs[0].text = translated_text
                                for run in paragraph.runs[1:]:
                                    run.text = ""
                                total_translated += 1

    doc.save(target_file)
    print(f"\n翻译完成！共翻译了 {total_translated} 个文本块。")
    print(f"文件已保存至: {target_file}")

if __name__ == "__main__":
    apply_translation()