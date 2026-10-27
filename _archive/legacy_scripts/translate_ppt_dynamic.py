import os
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from deep_translator import GoogleTranslator
from pptx.util import Pt

source_file = r"C:\Users\DMC86124042\Downloads\ASPD L3运维未尽事宜沟通_202605_CN_Revised3.pptx"
target_file = r"C:\Users\DMC86124042\Downloads\ASPD L3运维未尽事宜沟通_202605_CN_Revised3_JP.pptx"

translator = GoogleTranslator(source='zh-CN', target='ja')

memory_dict = {
    "双社长会": "ステコミ会議"
}

def translate_text(text):
    if not text.strip():
        return text
    
    # 替换带有特殊记忆的词
    for k, v in memory_dict.items():
        text = text.replace(k, v)
        
    try:
        translated = translator.translate(text)
        return translated
    except Exception as e:
        print(f"Translation Error for '{text}': {e}")
        return text

def apply_translation():
    prs = Presentation(source_file)
    total_translated = 0
    
    for slide_idx, slide in enumerate(prs.slides):
        print(f"Processing slide P{slide_idx + 1}...")
        for shape in slide.shapes:
            if shape.has_text_frame:
                # 判断是否为标题
                is_title = shape == slide.shapes.title if slide.shapes.title else False
                if is_title:
                    shape.text_frame.word_wrap = False
                
                # 开启文字自适应框大小
                try:
                    shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                except Exception:
                    pass

                for paragraph in shape.text_frame.paragraphs:
                    full_text = "".join(run.text for run in paragraph.runs).strip()
                    if full_text:
                        translated_text = translate_text(full_text)
                        if translated_text != full_text:
                            # 覆盖到第一个 run，清除其余的 run
                            if len(paragraph.runs) > 0:
                                font_size = paragraph.runs[0].font.size
                                paragraph.runs[0].text = translated_text
                                for run in paragraph.runs[1:]:
                                    run.text = ""
                                total_translated += 1

            # 如果形状有表格
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text_frame:
                            try:
                                cell.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                            except Exception:
                                pass
                            for paragraph in cell.text_frame.paragraphs:
                                full_text = "".join(run.text for run in paragraph.runs).strip()
                                if full_text:
                                    translated_text = translate_text(full_text)
                                    if translated_text != full_text:
                                        if len(paragraph.runs) > 0:
                                            paragraph.runs[0].text = translated_text
                                            for run in paragraph.runs[1:]:
                                                run.text = ""
                                            total_translated += 1

            # 处理群组类型内的形状 (Group shapes)
            if shape.shape_type == 6:
                for sub_shape in shape.shapes:
                    if sub_shape.has_text_frame:
                        try:
                            sub_shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                        except:
                            pass
                        for paragraph in sub_shape.text_frame.paragraphs:
                            full_text = "".join(run.text for run in paragraph.runs).strip()
                            if full_text:
                                translated_text = translate_text(full_text)
                                if translated_text != full_text:
                                    if len(paragraph.runs) > 0:
                                        paragraph.runs[0].text = translated_text
                                        for run in paragraph.runs[1:]:
                                            run.text = ""
                                        total_translated += 1

    prs.save(target_file)
    print(f"Translation completed. {total_translated} blocks translated. Saved to {target_file}")

if __name__ == "__main__":
    apply_translation()